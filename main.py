import os
import smtplib
import sqlite3
from datetime import datetime
from email.mime.multipart import MIMEMultipart
from email.mime.text import MIMEText

from fastapi import FastAPI, Form, Request
from fastapi.responses import HTMLResponse
from fastapi.staticfiles import StaticFiles
from fastapi.templating import Jinja2Templates
from openpyxl import Workbook, load_workbook

app = FastAPI()

app.mount("/static", StaticFiles(directory="static"), name="static")
templates = Jinja2Templates(directory="templates")

def init_db():
    conn = sqlite3.connect("database.db")
    cursor = conn.cursor()
    cursor.execute('''
        CREATE TABLE IF NOT EXISTS leads (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            nome_studio TEXT,
            nome_notaio TEXT,
            email TEXT,
            telefono TEXT,
            luogo TEXT,
            punteggio_totale INTEGER,
            area_prioritaria TEXT
        )
    ''')
    conn.commit()
    conn.close()

init_db()

# --- FUNZIONE SALVATAGGIO EXCEL ---
def salva_su_excel(nome_studio, nome_notaio, email, telefono, luogo, totale, profilo, area_prioritaria):
    file_excel = "risultati_quiz.xlsx"
    data_ora = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    
    try:
        if not os.path.exists(file_excel):
            wb = Workbook()
            ws = wb.active
            ws.title = "Leads Quiz"
            ws.append(["Data/Ora", "Nome Studio", "Nome Notaio", "Email", "Telefono", "Luogo", "Punteggio Totale", "Profilo", "Area Prioritaria"])
        else:
            wb = load_workbook(file_excel)
            ws = wb.active

        ws.append([data_ora, nome_studio, nome_notaio, email, telefono, luogo, totale, profilo, area_prioritaria])
        wb.save(file_excel)
    except Exception as e:
        print(f"Errore durante la scrittura su Excel: {e}")

# --- FUNZIONE NOTIFICA EMAIL GMAIL ---
def invia_email_notifica(nome_studio, nome_notaio, email_cliente, telefono, luogo, totale, profilo, area_prioritaria):
    mittente = "notaionextwki@gmail.com"
    password = "jlutonuzwyahquvj"  # <-- Sostituisci con la Password per le app di 16 lettere creata su Google
    destinatario = "notaionextwki@gmail.com"
    server_smtp = "smtp.gmail.com"
    porta_smtp = 587

    msg = MIMEMultipart()
    msg['From'] = mittente
    msg['To'] = destinatario
    msg['Subject'] = f"Nuovo Lead Quiz NotaioNext: {nome_studio} - {nome_notaio}"

    corpo = f"""
    Un nuovo utente ha completato il check-up su NotaioNext!

    DETTAGLI STUDIO:
    --------------------------------------------------
    Nome Studio: {nome_studio}
    Nome Notaio: {nome_notaio}
    Email Utente: {email_cliente}
    Telefono: {telefono}
    Luogo: {luogo}

    ESITO CHECK-UP:
    --------------------------------------------------
    Punteggio Totale: {totale} / 25
    Profilo Risultante: {profilo}
    Area Prioritaria da Intervenire: {area_prioritaria}

    Data compilazione: {datetime.now().strftime("%d/%m/%Y %H:%M")}
    --------------------------------------------------
    I dati sono stati salvati automaticamente anche nel file Excel 'risultati_quiz.xlsx' sul server IONOS.
    """
    msg.attach(MIMEText(corpo, 'plain', 'utf-8'))

    try:
        server = smtplib.SMTP(server_smtp, porta_smtp)
        server.starttls()
        server.login(mittente, password)
        server.send_message(msg)
        server.quit()
        print("Notifica email inviata con successo!")
    except Exception as e:
        print(f"Errore durante l'invio dell'e-mail: {e}")

# Mappatura dei Voucher d'offerta per ogni Area Prioritaria
VOUCHER_MAP = {
    "AI & Innovazione": {
        "titolo": "Corso AI & Innovazione Studio",
        "desc": "Un corso dedicato all'uso dell'Intelligenza Artificiale per automatizzare la gestione documentale e i processi dello Studio."
    },
    "Antiriciclaggio": {
        "titolo": "Corso Antiriciclaggio & Compliance",
        "desc": "Un corso dedicato all'antiriciclaggio per aggiornare e rafforzare competenze, procedure e organizzazione dello Studio."
    },
    "Organizzazione dello Studio": {
        "titolo": "Masterclass Cloud & Processi Operativi",
        "desc": "Un percorso per migrare in cloud e digitalizzare l'archivio e i flussi di lavoro operativi quotidiani."
    },
    "Comunicazione digitale": {
        "titolo": "Corso Digital Strategy & Marketing Notarile",
        "desc": "Strategie pratiche per ottimizzare la presenza online e la comunicazione istituzionale del tuo Studio."
    },
    "Controllo di gestione": {
        "titolo": "Masterclass Controllo di Gestione & KPI",
        "desc": "Implementa dashboard e indicatori chiave di performance per supportare le decisioni strategiche dello Studio."
    }
}

@app.get("/", response_class=HTMLResponse)
async def home(request: Request):
    return templates.TemplateResponse(request=request, name="index.html")

@app.post("/calcola", response_class=HTMLResponse)
async def calcola_risultati(
    request: Request,
    nome_studio: str = Form(...),
    nome_notaio: str = Form(...),
    email: str = Form(...),
    telefono: str = Form(...),
    luogo: str = Form(...),
    q1_ai: int = Form(...),
    q2_antiriciclaggio: int = Form(...),
    q3_organizzazione: int = Form(...),
    q4_comunicazione: int = Form(...),
    q5_controllo: int = Form(...)
):
    punteggi = {
        "AI & Innovazione": q1_ai,
        "Antiriciclaggio": q2_antiriciclaggio,
        "Organizzazione dello Studio": q3_organizzazione,
        "Comunicazione digitale": q4_comunicazione,
        "Controllo di gestione": q5_controllo
    }
    
    totale = sum(punteggi.values())
    
    if totale <= 10:
        profilo = "Studio Tradizionale"
    elif totale <= 15:
        profilo = "Studio in Evoluzione"
    elif totale <= 20:
        profilo = "Studio Digitale"
    else:
        profilo = "Studio Evoluto"
        
    area_prioritaria = min(punteggi, key=punteggi.get)
    punti_area = punteggi[area_prioritaria]
    voucher_info = VOUCHER_MAP.get(area_prioritaria, VOUCHER_MAP["Antiriciclaggio"])

    percentuale_pos = int(((totale - 5) / 20) * 100)

    # 1. Salvataggio in SQLite
    conn = sqlite3.connect("database.db")
    cursor = conn.cursor()
    cursor.execute('''
        INSERT INTO leads (nome_studio, nome_notaio, email, telefono, luogo, punteggio_totale, area_prioritaria)
        VALUES (?, ?, ?, ?, ?, ?, ?)
    ''', (nome_studio, nome_notaio, email, telefono, luogo, totale, area_prioritaria))
    conn.commit()
    conn.close()

    # 2. Salvataggio automatico nel file Excel (risultati_quiz.xlsx)
    salva_su_excel(nome_studio, nome_notaio, email, telefono, luogo, totale, profilo, area_prioritaria)

    # 3. Invia notifica e-mail via Gmail
    invia_email_notifica(nome_studio, nome_notaio, email, telefono, luogo, totale, profilo, area_prioritaria)

    return templates.TemplateResponse(
        request=request,
        name="risultati.html",
        context={
            "nome_studio": nome_studio,
            "nome_notaio": nome_notaio,
            "totale": totale,
            "profilo": profilo,
            "punteggi": punteggi,
            "area_prioritaria": area_prioritaria,
            "punti_area": punti_area,
            "voucher_info": voucher_info,
            "percentuale_pos": percentuale_pos
        }
    )
